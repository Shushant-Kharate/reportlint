import io
import zipfile
from pathlib import Path
import pytest
from pydantic import ValidationError
from docx import Document
from docx.shared import Pt
from fastapi import UploadFile, HTTPException
from app.models.document_model import DocumentModel, ParagraphModel, RunModel, SectionPropertiesModel, Confidence
from app.models.rule_model import Rule, RuleSet, RequiredSectionRule
from app.models.result_model import Violation
from app.compliance.engine import ComplianceEngine
from app.compliance.scoring import compute_scores
from app.ooxml.docx_loader import DocxPackage, InvalidDocxError
from app.ooxml.structure_extractor import build_document_model
from app import storage
from tests.test_stage2_api import client


def rule(kind='FONT_SIZE', expected=None, **kwargs):
    return Rule(id='test', type=kind, scope='DOCUMENT', expected_value=expected or {'size_pt':12}, severity='ERROR', **kwargs)


def test_unchecked_categories_do_not_inflate_score():
    violation=Violation(rule_id='x',severity='ERROR',category='TYPOGRAPHY',message='Wrong',expected={})
    score, categories=compute_scores([violation], {'TYPOGRAPHY':1})
    assert score == 0
    assert all(c.score == 0 for c in categories if c.checks_performed == 0)
    assert compute_scores([], {})[0] == 0


def test_known_size_is_checked_even_when_font_is_unknown():
    doc=DocumentModel(source_filename='x',paragraphs=[ParagraphModel(index=0,text='Text',runs=[RunModel(text='Text',font_size_pt=14,confidence=Confidence.UNKNOWN)])])
    result=ComplianceEngine(RuleSet(template_source_filename='x',typography_rules=[rule()])).run(doc)
    assert result.total_errors == 1
    assert result.overall_score == 0


def test_unknown_value_is_not_counted_as_a_pass():
    doc=DocumentModel(source_filename='x',paragraphs=[ParagraphModel(index=0,text='Text',runs=[RunModel(text='Text',font_size_pt=None,confidence=Confidence.EXPLICIT)])])
    result=ComplianceEngine(RuleSet(template_source_filename='x',typography_rules=[rule()])).run(doc)
    assert result.total_checks_passed == 0


def test_optional_and_info_sections_are_not_enforced():
    rs=RuleSet(template_source_filename='x',structure_rules=[RequiredSectionRule(canonical_name='optional',required=False,order_index=0),RequiredSectionRule(canonical_name='info',severity='INFO',order_index=1)])
    result=ComplianceEngine(rs).run(DocumentModel(source_filename='x'))
    assert result.total_errors == 0
    assert result.total_checks_passed == 0


def test_page_checks_cover_later_sections():
    doc=DocumentModel(source_filename='x',sections=[SectionPropertiesModel(page_width_pt=595,page_height_pt=842),SectionPropertiesModel(page_width_pt=612,page_height_pt=792)])
    rs=RuleSet(template_source_filename='x',page_rules=[rule('PAGE_SIZE',{'width_pt':595,'height_pt':842})])
    result=ComplianceEngine(rs).run(doc)
    assert result.total_errors == 1
    assert result.total_checks_passed == 1
    assert result.violations[0].location.section == 'Document section 2'


def test_exact_line_spacing_is_enforced():
    doc=DocumentModel(source_filename='x',paragraphs=[ParagraphModel(index=0,text='Body',line_spacing_rule='auto',line_spacing=1.5)])
    rs=RuleSet(template_source_filename='x',paragraph_rules=[rule('LINE_SPACING_EXACT',{'pt':18})])
    assert ComplianceEngine(rs).run(doc).total_errors == 1


def test_exact_spacing_does_not_get_overwritten_by_inherited_multiple(tmp_path):
    d=Document();d.styles['Normal'].paragraph_format.line_spacing=1.5
    p=d.add_paragraph('Exact spacing');p.paragraph_format.line_spacing=Pt(18)
    path=tmp_path/'exact.docx';d.save(path)
    doc=build_document_model(DocxPackage.load(str(path)),'x')
    assert doc.paragraphs[0].line_spacing is None
    assert doc.paragraphs[0].line_spacing_exact_pt == 18
    assert doc.paragraphs[0].line_spacing_rule == 'exact'


@pytest.mark.parametrize('value', ['twelve',None,float('nan'),-1,True])
def test_invalid_numeric_rules_are_rejected(value):
    with pytest.raises(ValidationError):
        rule(expected={'size_pt':value})


def test_unsupported_enforced_rule_rejected():
    with pytest.raises(ValidationError):
        rule('CAPTION_PRESENCE',{'present':True})


def test_malformed_document_returns_400(client):
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w') as z:
        z.writestr('word/document.xml','<root/>')
    response=client.post('/api/templates',files={'file':('bad.docx',stream.getvalue())})
    assert response.status_code == 400


def test_xml_entities_are_rejected(tmp_path):
    path=tmp_path/'entity.docx'
    with zipfile.ZipFile(path,'w') as z:
        z.writestr('word/document.xml','<!DOCTYPE x [<!ENTITY x "payload">]><x>&x;</x>')
    with pytest.raises(InvalidDocxError):
        DocxPackage.load(str(path))


def test_storage_cannot_escape_root(tmp_path,monkeypatch):
    monkeypatch.setattr(storage,'TEMPLATES_DIR',tmp_path/'templates')
    assert storage.get_template('../outside') is None
    assert not storage.delete_template('../outside')


def test_bad_template_file_does_not_break_listing(tmp_path,monkeypatch):
    monkeypatch.setattr(storage,'TEMPLATES_DIR',tmp_path)
    (tmp_path/'bad.json').write_text('{broken')
    assert storage.list_templates() == []


def test_saving_published_template_returns_it_to_draft(client):
    with open('tests/fixtures/mini_project_1a_format.docx','rb') as f:
        t=client.post('/api/templates',files={'file':('format.docx',f)}).json()
    client.post(f"/api/templates/{t['id']}/publish")
    r=client.put(f"/api/templates/{t['id']}/rules",json=t['ruleset'])
    assert r.status_code == 200
    assert r.json()['status'] == 'DRAFT'


def test_upload_size_limit_is_bounded():
    import asyncio
    from app.uploads import read_limited
    upload=UploadFile(file=io.BytesIO(b'12345'),filename='test.docx')
    with pytest.raises(HTTPException) as exc:
        asyncio.run(read_limited(upload,4))
    assert exc.value.status_code == 413


def test_upload_temp_directory_is_cleaned(monkeypatch):
    from app import uploads
    original=uploads.tempfile.TemporaryDirectory
    paths=[]
    def record(*args,**kwargs):
        obj=original(*args,**kwargs);paths.append(Path(obj.name));return obj
    monkeypatch.setattr(uploads.tempfile,'TemporaryDirectory',record)
    with pytest.raises(HTTPException):
        uploads._parse(b'invalid','test.docx')
    assert paths and not paths[0].exists()

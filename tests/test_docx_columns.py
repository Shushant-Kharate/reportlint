from io import BytesIO

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

from app.builtin_formats import builtin
from app.checking.columns import column_count
from app.checking.service import check_bytes
from app.ooxml.constants import qn

TEXT = ('This paragraph explains the experiment and its measured results in ordinary sentences. '
        'The document discusses a repeatable method and evaluates the findings carefully. ') * 5


def set_columns(section,count):
    cols = section._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols');section._sectPr.append(cols)
    cols.set(qn('w:num'),str(count))
    return cols


def encode(document):
    stream=BytesIO();document.save(stream);return stream.getvalue()


def example(count=None,style='Body Text'):
    doc=Document()
    doc.sections[0].page_width=Inches(8.5)
    doc.sections[0].page_height=Inches(11)
    doc.styles['Normal'].font.name='Times New Roman'
    doc.styles['Body Text'].font.name='Times New Roman'
    doc.styles['Body Text'].font.size=Pt(10)
    if count is not None:
        set_columns(doc.sections[0],count)
    for _ in range(2):
        doc.add_paragraph(TEXT,style)
    return doc


def check(document):
    return check_bytes(builtin('builtin-ieee-conference-letter'),encode(document))


def column_items(result):
    return [i for i in result.items if i.expected and i.expected.get('kind')=='body_columns']


def test_matching_paper_and_font_do_not_hide_single_column_body():
    result=check(example())
    assert any(i.status=='PASS' and i.code=='STORED_SETTING_MATCH' for i in result.items)
    assert any(i.status=='PASS' and i.code=='BODY_RUN_PASS' for i in result.items)
    item=column_items(result)[0]
    assert result.outcome=='FAIL' and item.status=='FAIL'
    assert item.actual['observed_columns']==1 and item.expected['count']==2
    assert item.section_index==0 and item.actual['sample_paragraph_indices']


def test_changing_only_body_columns_removes_the_failure():
    doc=example(1)
    before=check(doc)
    set_columns(doc.sections[0],2)
    after=check(doc)
    assert before.counts['FAIL']==1 and after.counts['FAIL']==0
    assert column_items(after)[0].status=='PASS'


def test_normal_style_prose_can_establish_layout_without_font_role_guessing():
    result=check(example(1,'Normal'))
    item=column_items(result)[0]
    assert item.status=='FAIL' and 'prose' in item.actual['scope_basis']
    assert any(i.code=='INCOMPLETE_BODY_SCOPE' for i in result.items)


def test_one_column_authors_and_full_width_caption_are_not_failed():
    doc=Document()
    set_columns(doc.sections[0],1)
    doc.add_paragraph('Authors and affiliations','Title')
    section=doc.add_section(WD_SECTION.CONTINUOUS);set_columns(section,2)
    doc.add_paragraph(TEXT,'Body Text');doc.add_paragraph(TEXT,'Body Text')
    section=doc.add_section(WD_SECTION.CONTINUOUS);set_columns(section,1)
    doc.add_paragraph('Figure 1. '+TEXT,'Caption')
    doc.add_table(rows=1,cols=1).cell(0,0).text=TEXT
    items=column_items(check(doc))
    assert [i.status for i in items]==['NOT_CHECKED','PASS','NOT_CHECKED']


def test_missing_column_element_uses_single_column_default():
    doc=example()
    sect=doc.sections[0]._sectPr
    sect.remove(sect.find(qn('w:cols')))
    assert column_count(sect)[0]==1
    assert column_items(check(doc))[0].status=='FAIL'


def test_unequal_columns_use_valid_explicit_definitions():
    doc=example(2);cols=set_columns(doc.sections[0],2)
    cols.set(qn('w:equalWidth'),'0')
    for width in ['3000','4000']:
        child=OxmlElement('w:col');child.set(qn('w:w'),width);cols.append(child)
    assert column_items(check(doc))[0].status=='PASS'
    cols.set(qn('w:num'),'3')
    assert column_items(check(doc))[0].status=='NOT_CHECKED'


@pytest.mark.parametrize('value',['0','-1','bogus','99'])
def test_malformed_column_count_is_not_called_a_violation(value):
    doc=example(2);set_columns(doc.sections[0],value)
    assert column_items(check(doc))[0].status=='NOT_CHECKED'


def test_ambiguous_section_history_is_left_for_review():
    doc=example(1)
    history=OxmlElement('w:sectPrChange');history.append(OxmlElement('w:sectPr'))
    doc.sections[0]._sectPr.append(history)
    items=column_items(check(doc))
    assert len(items)==1 and items[0].status=='NOT_CHECKED'


def test_short_cover_is_not_forced_into_two_columns():
    doc=Document();doc.add_paragraph('A report cover and author names')
    assert all(i.status=='NOT_CHECKED' for i in column_items(check(doc)))

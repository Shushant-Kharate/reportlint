from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt
from fastapi.testclient import TestClient
from pydantic import ValidationError
import pytest

from app.checking.role_review import ReportRoleReview, RoleDecision, preview_bytes
from app.checking.service import check_bytes
from app.main import app
from app.ooxml.constants import qn
from app.review.api import get_repository
from app.review.service import ReviewError
from tests.test_semantic_check import revision
from tests.test_template_analysis import encoded


def manual_report():
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.5
    doc.add_paragraph("A fictional report cover")
    for title in ["Introduction", "System Design", "References"]:
        doc.add_paragraph(title).runs[0].bold = True
        doc.add_paragraph("This is a complete fictional paragraph with sufficient words to suggest ordinary body prose.")
    return doc


def review_all(preview):
    return ReportRoleReview(report_sha256=preview.report_sha256, snapshot_sha256=preview.snapshot_sha256,
        decisions=[RoleDecision(paragraph_index=p.paragraph_index, source_path=p.source_path,
                    role="EXCLUDE" if p.paragraph_index == 0 else p.suggested_role,
                    chapter_index=p.suggested_chapter_index,
                    reason="Synthetic reviewer verified this paragraph against its location in the report.")
                   for p in preview.paragraphs if p.reviewable])


def test_preview_suggests_but_does_not_approve(revision):
    data = encoded(manual_report())
    before = check_bytes(revision, data)
    preview = preview_bytes(revision, data)
    assert preview.counts == {"UNKNOWN": 7}
    assert [p.suggested_role for p in preview.paragraphs] == [None, "CHAPTER", "BODY", "CHAPTER", "BODY", "CHAPTER", "BODY"]
    assert all(p.review_status == "PROPOSED" for p in preview.paragraphs)
    assert check_bytes(revision, data).model_dump() == before.model_dump()
    assert any(i.code == "NO_TRUSTED_CHAPTER_HEADINGS" or i.code == "AMBIGUOUS_CHAPTER" for i in before.items)


def test_manual_assignments_enable_body_and_chapter_checks(revision):
    data = encoded(manual_report())
    preview = preview_bytes(revision, data)
    review = review_all(preview)
    result = check_bytes(revision, data, review)
    assert len([i for i in result.items if i.code == "CHAPTER_PRESENT"]) == 3
    body = [i for i in result.items if i.code.startswith("BODY_")]
    assert len(body) == 9 and all(i.status == "PASS" for i in body)
    assert not any(i.code == "INCOMPLETE_BODY_SCOPE" for i in result.items)
    assert result.role_review_sha256
    assert result.role_decisions == review.decisions
    assert any(i.code == "REVIEWED_ROLE_EXCLUSION" and i.status == "OUT_OF_SCOPE" for i in result.items)
    assert check_bytes(revision, data, review).model_dump() == result.model_dump()


def test_manual_body_still_finds_wrong_size(revision):
    doc = manual_report()
    doc.paragraphs[2].runs[0].font.size = Pt(10)
    data = encoded(doc)
    result = check_bytes(revision, data, review_all(preview_bytes(revision, data)))
    errors = [i for i in result.items if i.code == "BODY_RUN_FAIL"]
    assert len(errors) == 1 and errors[0].paragraph_index == 2


@pytest.mark.parametrize("kind", ["report", "snapshot", "path", "index", "duplicate", "chapter"])
def test_stale_or_invalid_role_reviews_are_rejected(revision, kind):
    data = encoded(manual_report())
    review = review_all(preview_bytes(revision, data))
    if kind == "report":
        review.report_sha256 = "0" * 64
    elif kind == "snapshot":
        review.snapshot_sha256 = "0" * 64
    elif kind == "path":
        review.decisions[0].source_path = "/other"
    elif kind == "index":
        review.decisions[0].paragraph_index = 999
    elif kind == "duplicate":
        review.decisions.append(review.decisions[0])
    else:
        review.decisions[1].chapter_index = 999
    with pytest.raises(ReviewError) as error:
        check_bytes(revision, data, review)
    assert error.value.code == ("STALE_ROLE_REVIEW" if kind in {"report", "snapshot"} else "INVALID_CHAPTER_ASSIGNMENT" if kind == "chapter" else "INVALID_ROLE_TARGET")


@pytest.mark.parametrize("kind", ["table", "field", "style_cycle", "drawing"])
def test_unsafe_roles_cannot_be_overridden(revision, kind):
    doc = Document()
    if kind == "table":
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    else:
        p = doc.add_paragraph()
    p.text = "Introduction"
    if kind == "field":
        p._p.append(OxmlElement("w:fldSimple"))
    elif kind == "style_cycle":
        doc.styles["Normal"].base_style = doc.styles["Normal"]
    elif kind == "drawing":
        p.runs[0]._r.append(OxmlElement("w:drawing"))
    data = encoded(doc)
    preview = preview_bytes(revision, data)
    target = preview.paragraphs[0]
    assert not target.reviewable
    review = ReportRoleReview(report_sha256=preview.report_sha256, snapshot_sha256=preview.snapshot_sha256,
        decisions=[RoleDecision(paragraph_index=target.paragraph_index, source_path=target.source_path, role="BODY", reason="Cannot bypass the parser's uncertainty.")])
    with pytest.raises(ReviewError) as error:
        check_bytes(revision, data, review)
    assert error.value.code == "UNSUPPORTED_ROLE_OVERRIDE"


def test_unreviewed_paragraphs_remain_unchecked(revision):
    data = encoded(manual_report())
    review = review_all(preview_bytes(revision, data))
    review.decisions.pop()
    assert any(i.code == "INCOMPLETE_BODY_SCOPE" for i in check_bytes(revision, data, review).items)


def test_repeated_manual_title_requires_individual_review(revision):
    doc = manual_report()
    doc.add_paragraph("Introduction")
    data = encoded(doc)
    preview = preview_bytes(revision, data)
    review = review_all(preview)
    assert len([p for p in preview.paragraphs if p.suggested_chapter_index == 0]) == 2
    assert any(i.code == "DUPLICATE_CHAPTER" for i in check_bytes(revision, data, review).items)
    review.decisions[-1] = RoleDecision(paragraph_index=7, source_path=preview.paragraphs[7].source_path, role="EXCLUDE", reason="Reviewer confirms this is a manually typed contents mention.")
    assert not any(i.code in {"DUPLICATE_CHAPTER", "AMBIGUOUS_CHAPTER"} for i in check_bytes(revision, data, review).items)


def test_explicit_mapping_supports_different_title_without_changing_template(revision):
    original = revision.model_dump()
    doc = manual_report()
    doc.paragraphs[1].text = "Chapter One - Introductory Context"
    data = encoded(doc)
    preview = preview_bytes(revision, data)
    assert preview.paragraphs[1].suggested_role is None
    decision = RoleDecision(paragraph_index=1, source_path=preview.paragraphs[1].source_path, role="CHAPTER", chapter_index=0, reason="Reviewer verified this report paragraph introduces the required chapter.")
    review = ReportRoleReview(report_sha256=preview.report_sha256, snapshot_sha256=preview.snapshot_sha256, decisions=[decision])
    assert any(i.code == "CHAPTER_PRESENT" and i.paragraph_index == 1 for i in check_bytes(revision, data, review).items)
    assert revision.model_dump() == original


def test_tabs_and_breaks_are_preserved_in_preview(revision):
    doc = Document()
    doc.add_paragraph("System\tDesign\nExample")
    p = preview_bytes(revision, encoded(doc)).paragraphs[0]
    assert p.excerpt == "System\tDesign\nExample" and p.suggested_role is None


def test_api_preview_and_reviewed_check(revision, tmp_path):
    class ReadOnlyRepository:
        def get(self, template_id, revision_id):
            assert (template_id, revision_id) == (revision.template_id, revision.revision_id)
            return revision
    app.dependency_overrides[get_repository] = lambda: ReadOnlyRepository()
    try:
        with TestClient(app) as client:
            base = f"/api/v2/templates/{revision.template_id}/revisions/{revision.revision_id}"
            data = encoded(manual_report())
            preview = client.post(base + "/role-preview", files={"file": ("manual.docx", data)})
            assert preview.status_code == 200
            review = review_all(preview_bytes(revision, data))
            response = client.post(base + "/check", files={"file": ("manual.docx", data)}, data={"role_review": review.model_dump_json()})
            assert response.status_code == 200 and response.json()["role_review_sha256"]
            assert client.post(base + "/check", files={"file": ("manual.docx", data)}, data={"role_review": '{"decisions":[]}'}).status_code == 422
            review.report_sha256 = "0" * 64
            stale = client.post(base + "/check", files={"file": ("manual.docx", data)}, data={"role_review": review.model_dump_json()})
            assert stale.status_code == 409 and stale.json()["detail"]["code"] == "STALE_ROLE_REVIEW"
    finally:
        app.dependency_overrides.pop(get_repository, None)


@pytest.mark.parametrize("changes", [
    {"role": "CHAPTER"}, {"chapter_index": 0}, {"reason": ""},
    {"paragraph_index": -1}, {"paragraph_index": "0"}, {"execute": "arbitrary operation"},
])
def test_role_decision_schema_rejects_invalid_input(changes):
    value = {"paragraph_index": 0, "source_path": "/w:document/w:body/w:p[1]", "role": "BODY", "reason": "Reviewed prose paragraph."}
    value.update(changes)
    with pytest.raises(ValidationError):
        RoleDecision.model_validate(value)


def test_true_default_style_inheritance_is_retained(revision):
    doc = manual_report()
    doc.styles["Normal"]._element.set(qn("w:default"), "true")
    data = encoded(doc)
    result = check_bytes(revision, data, review_all(preview_bytes(revision, data)))
    assert all(i.status == "PASS" for i in result.items if i.code.startswith("BODY_"))


def test_malformed_style_reference_is_not_reviewable(revision):
    doc = manual_report()
    style = OxmlElement("w:pStyle")
    doc.paragraphs[2]._p.get_or_add_pPr().append(style)
    assert not preview_bytes(revision, encoded(doc)).paragraphs[2].reviewable

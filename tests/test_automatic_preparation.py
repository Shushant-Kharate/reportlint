from io import BytesIO
from copy import deepcopy
import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient
from app.main import app
from app.review.api import get_repository
from app.review.repository import RevisionRepository
from app.builtin_formats import CatalogRepository,DIRECTORY,builtin
from app.analysis.plain_language import interpret
from app.review.service import snapshot_hash
from tests.test_template_analysis import encoded


@pytest.fixture
def client(tmp_path):
    repository=CatalogRepository(RevisionRepository(tmp_path/'automatic.sqlite3'))
    app.dependency_overrides[get_repository]=lambda:repository
    with TestClient(app) as c:yield c
    app.dependency_overrides.pop(get_repository,None)


def upload(client,lines,sample=False):
    doc=Document()
    for line in lines:doc.add_paragraph(line)
    if sample:
        doc.styles['Body Text'].font.name='Arial';doc.styles['Body Text'].font.size=Pt(14)
        doc.add_paragraph('This sample describes a fictional process and its measured results. '*30,'Body Text')
    response=client.post('/api/v2/templates/prepare',files={'file':('new-format.docx',encoded(doc))})
    assert response.status_code==200,response.text
    return response.json()


def path(result):
    r=result['revision']
    return f"/api/v2/templates/{r['template_id']}/revisions/{r['revision_id']}"


def rules(result,kind):
    return [r for r in (result['revision']['publication'] or {}).get('rules',[]) if r['value']['kind']==kind]


def test_known_ieee_upload_reuses_nineteen_rules_without_review_or_duplicate(client):
    result=client.post('/api/v2/templates/prepare',files={'file':('renamed-format.docx',(DIRECTORY/'ieee-conference-letter.docx').read_bytes())}).json()
    assert result['reused'] and result['ready'] and result['rule_count']==19
    assert result['revision']['revision_id']=='builtin-ieee-conference-letter-v2'
    assert len(client.get('/api/v2/templates').json())==2


def test_clear_requirements_publish_without_approval_reasons(client):
    result=upload(client,['Use A4 paper.','Body text must use Times New Roman at 12 pt with 1.5 line spacing in two columns and justified alignment.','All margins must be 1 inch.'])
    assert result['ready'] and result['revision']['status']=='PUBLISHED'
    assert not result['questions']
    assert rules(result,'font_size')[0]['value']['expected_pt']==12
    assert rules(result,'body_columns')[0]['value']['count']==2
    assert len(rules(result,'margin'))==4
    assert result['revision']['snapshot_sha256']==snapshot_hash(__import__('app.review.models',fromlist=['ReviewRevision']).ReviewRevision.model_validate(result['revision']))


def test_unknown_clause_remains_a_question_and_unchecked_result(client):
    result=upload(client,['Body text must use Times New Roman at 12 pt.','All lettering must be embossed in gold.'])
    assert result['ready'] and len(result['questions'])==1
    report=Document();report.styles['Body Text'].font.name='Times New Roman';report.styles['Body Text'].font.size=Pt(12)
    report.add_paragraph('A fictional body paragraph for the report.','Body Text')
    response=client.post(path(result)+'/check',files={'file':('report.docx',encoded(report))}).json()
    assert any(i['code']=='REQUIREMENT_MANUAL' and i['status']=='NOT_CHECKED' for i in response['items'])
    assert not any(i['code']=='REQUIREMENT_COVERED' for i in response['items'])


def test_explicit_values_override_conflicting_sample_appearance(client):
    result=upload(client,['Body text must use Times New Roman at 10 pt.'],sample=True)
    assert result['ready']
    assert [r['value']['font'] for r in rules(result,'font_family')]==['Times New Roman']
    assert [r['value']['expected_pt'] for r in rules(result,'font_size')]==[10]


def test_alternatives_do_not_silently_choose_a_font_or_sample(client):
    result=upload(client,['Use A4 paper.','Body text must use Arial or Times New Roman at 12 pt.'],sample=True)
    assert result['ready'] and result['questions']
    assert not rules(result,'font_family') and not rules(result,'font_size')


def test_conflicting_explicit_sizes_remain_unresolved(client):
    result=upload(client,['Body text must use Times New Roman at 10 pt.','Body text must use 12 pt.'])
    assert result['ready'] and not rules(result,'font_size')
    assert len(result['questions'])==2
    updated=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'evidence_id':result['questions'][0]['evidence_id'],'explanation':'Body text must use Times New Roman at 12 pt.'})
    assert updated.status_code==200,updated.text
    assert not updated.json()['questions']


def test_clarification_replaces_ambiguity_and_preserves_old_snapshot(client):
    result=upload(client,['Use A4 paper.','Body text must use Arial or Times New Roman at 12 pt.'])
    old=deepcopy(result['revision'])
    response=client.post(path(result)+'/clarify',json={'expected_version':old['version'],
        'evidence_id':result['questions'][0]['evidence_id'],
        'explanation':'Body text must use Arial at 12 pt with 1.5 line spacing.'})
    assert response.status_code==200,response.text
    revised=response.json()
    assert revised['ready'] and not revised['questions']
    assert rules(revised,'font_family')[0]['value']['font']=='Arial'
    assert revised['revision']['revision_number']==2
    assert client.get(path(result)).json()==old
    assert any(e['part']=='user-clarification' for e in revised['revision']['analysis']['evidence'])


def test_explanation_changes_actual_report_comparison(client):
    result=upload(client,['Body text must use Times New Roman at 10 pt.'])
    doc=Document();doc.styles['Body Text'].font.name='Times New Roman';doc.styles['Body Text'].font.size=Pt(12)
    doc.add_paragraph('This fictional paragraph describes a test.','Body Text')
    before=client.post(path(result)+'/check',files={'file':('report.docx',encoded(doc))}).json()
    assert any(i['status']=='FAIL' and i['expected']['kind']=='font_size' for i in before['items'] if i.get('expected'))
    updated=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'explanation':'Body text must use 12 pt.'}).json()
    after=client.post(path(updated)+'/check',files={'file':('report.docx',encoded(doc))}).json()
    assert not any(i['status']=='FAIL' and i['expected']['kind']=='font_size' for i in after['items'] if i.get('expected'))


@pytest.mark.parametrize('explanation',['Use any beautiful font.','Body text must use 10 pt and 12 pt.','Body text must use Arial or Calibri.'])
def test_rejected_explanations_do_not_create_extra_revisions(client,explanation):
    result=upload(client,['Use A4 paper.'])
    response=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],'explanation':explanation})
    assert response.status_code==422
    assert len(client.get(path(result).rsplit('/',1)[0]).json())==1


def test_stale_explanation_is_rejected(client):
    result=upload(client,['Use A4 paper.'])
    assert client.post(path(result)+'/clarify',json={'expected_version':0,'explanation':'Body text must use 12 pt.'}).status_code==409
    first=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],'explanation':'Body text must use 12 pt.'})
    assert first.status_code==200
    second=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],'explanation':'Body text must use 14 pt.'})
    assert second.status_code==409


def test_repeated_custom_upload_reuses_published_revision(client):
    one=upload(client,['Body text must use Times New Roman at 12 pt.'])
    # DOCX metadata may vary; repeat the exact file bytes to test identity.
    doc=Document();doc.add_paragraph('Body text must use Arial at 11 pt.');data=encoded(doc)
    a=client.post('/api/v2/templates/prepare',files={'file':('a.docx',data)}).json()
    b=client.post('/api/v2/templates/prepare',files={'file':('b.docx',data)}).json()
    assert b['reused'] and a['revision']['revision_id']==b['revision']['revision_id']


def test_required_chapters_are_explicit_not_guessed_from_topic_table(client):
    result=upload(client,['Use A4 paper.','Required chapters: Introduction, Methods, Results, Conclusion.'])
    profile=result['revision']['publication']['profile']
    assert [c['name'] for c in profile['chapters']]==['Introduction','Methods','Results','Conclusion']
    assert all(c['required'] for c in profile['chapters'])


def test_scans_and_unsupported_explanations_are_not_passes(client):
    result=upload(client,['This format must look professional and use a blue leather binding.'])
    assert result['questions']


def test_targeted_short_answer_uses_the_instruction_role(client):
    result=upload(client,['Use A4 paper.','Body text must use Arial or Times New Roman.'])
    response=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'evidence_id':result['questions'][0]['evidence_id'],'explanation':'Use Arial at 12 pt.'})
    assert response.status_code==200,response.text
    updated=response.json()
    assert not updated['questions'] and rules(updated,'font_size')[0]['value']['expected_pt']==12


def test_named_style_baseline_is_distinct_from_a_dominant_caption_sample(client):
    doc=Document();doc.styles['Body Text'].font.name='Georgia';doc.styles['Body Text'].font.size=Pt(12)
    doc.add_paragraph('This fictional paragraph describes a sample and its observations. '*30,'Body Text')
    for _ in range(10):doc.add_paragraph('Small caption text.','Caption').runs[0].font.size=Pt(8)
    result=client.post('/api/v2/templates/prepare',files={'file':('example-format.docx',encoded(doc))}).json()
    assert rules(result,'font_family')[0]['value']['font']=='Georgia'
    assert rules(result,'font_size')[0]['value']['expected_pt']==12


def test_table_row_can_define_a_body_font_and_size(client):
    doc=Document();table=doc.add_table(rows=1,cols=3)
    for c,t in zip(table.rows[0].cells,['Body font','Times New Roman','12 pt']):c.text=t
    response=client.post('/api/v2/templates/prepare',files={'file':('table-format.docx',encoded(doc))})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['ready'] and rules(result,'font_size')[0]['value']['expected_pt']==12


def test_image_requirements_stay_visible_until_the_user_explains_them(client):
    import struct,zlib
    def chunk(name,data):return struct.pack('>I',len(data))+name+data+struct.pack('>I',zlib.crc32(name+data)&0xffffffff)
    image=BytesIO(b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR',struct.pack('>IIBBBBB',2,2,8,2,0,0,0))+chunk(b'IDAT',zlib.compress((b'\0'+b'\xff'*6)*2))+chunk(b'IEND',b''))
    doc=Document();doc.add_paragraph('Use A4 paper.');doc.add_picture(image)
    result=client.post('/api/v2/templates/prepare',files={'file':('pictured-format.docx',encoded(doc))}).json()
    question=next(q for q in result['questions'] if q['evidence_id']=='automatic-visual-content-question')
    assert question and result['ready']
    response=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'evidence_id':question['evidence_id'],'explanation':'This image is decorative.'})
    assert response.status_code==200,response.text
    assert not response.json()['questions']


def test_covered_cannot_hide_an_unsupported_instruction(client):
    doc=Document();doc.add_paragraph('Body text must use Times New Roman at 12 pt.');doc.add_paragraph('All lettering must be embossed in gold.')
    r=client.post('/api/v2/templates',files={'file':('manual.docx',encoded(doc))}).json()
    p=f"/api/v2/templates/{r['template_id']}/revisions/{r['revision_id']}"
    edited=client.patch(p,json={'expected_version':0,
        'candidates':[{'candidate_id':c['candidate_id'],'action':'APPROVE' if c['origin']=='EXPLICIT_PROSE' else 'REJECT','reason':'Synthetic review decision.'} for c in r['analysis']['candidates']],
        'ledger':[{'evidence_id':e['evidence_id'],'action':'COVERED','reason':'Attempting to claim unsupported coverage.'} for e in r['analysis']['requirement_ledger']]}).json()
    response=client.post(p+'/publish',json={'expected_version':edited['version']})
    assert response.status_code==409
    assert any(b['code']=='UNPROVEN_COVERAGE' for b in response.json()['detail']['blockers'])


def test_uploaded_known_format_remembers_private_clarifications(client):
    data=(DIRECTORY/'ieee-conference-letter.docx').read_bytes()
    result=client.post('/api/v2/templates/prepare',files={'file':('ieee.docx',data)}).json()
    modified=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'explanation':'Body text must use 12 pt.'})
    assert modified.status_code==200,modified.text
    repeated=client.post('/api/v2/templates/prepare',files={'file':('renamed.docx',data)}).json()
    assert repeated['reused']
    assert repeated['revision']['revision_id']==modified.json()['revision']['revision_id']
    assert rules(repeated,'font_size')[0]['value']['expected_pt']==12


@pytest.mark.parametrize('text',['Body text must use 12 pt on the first page.',
    'Body text in the appendix must use 12 pt.', 'Numbered paragraphs in the body must use 12 pt.',
    'No page numbers on the first page.'])
def test_region_specific_requirements_are_not_promoted_to_global_rules(text):
    parsed=interpret(text)
    assert parsed.ambiguity and not parsed.values


def test_paper_question_supplies_context_for_short_answer(client):
    result=upload(client,['Body text must use 12 pt.','Use A4 or Letter paper.'])
    response=client.post(path(result)+'/clarify',json={'expected_version':result['revision']['version'],
        'evidence_id':result['questions'][0]['evidence_id'],'explanation':'Use A4.'})
    assert response.status_code==200,response.text
    assert rules(response.json(),'page_size')[0]['value']['width_pt']==pytest.approx(210*72/25.4)


@pytest.mark.parametrize('text,kind',[('Body text should be double-spaced.','line_spacing_multiple'),
    ('No page numbers.','paper_no_page_numbers'),('Main headings must use uppercase Roman numerals.','paper_heading_numbering'),
    ('Abstract text must use 9 pt.','paper_abstract_size'),('Paper title must use 24 pt.','paper_title_size')])
def test_plain_language_forms(text,kind):
    parsed=interpret(text)
    assert parsed.complete and not parsed.ambiguity,parsed.remaining
    assert any(v.kind==kind for _,v in parsed.values)

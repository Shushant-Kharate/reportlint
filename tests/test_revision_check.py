from hashlib import sha256

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.shared import Pt
from fastapi.testclient import TestClient
import pytest

from app.analysis.service import analyze_bytes
from app.checking.service import check_bytes
from app.main import app
from app.ooxml.constants import qn
from app.ooxml.docx_loader import InvalidDocxError
from app.review.api import get_repository
from app.review.models import CandidateDecision, ReviewPatch
from app.review.repository import RevisionRepository
from app.review.service import apply_review, publish, ReviewError, snapshot_hash
from tests.test_template_analysis import encoded, complex_format
from tests.test_revision_review import complete_patch


@pytest.fixture
def repository(tmp_path):
    return RevisionRepository(tmp_path / "review.sqlite3")


@pytest.fixture
def published(repository):
    # Approve observed page settings of a fictional plain document only.
    draft = repository.create("Page settings fixture", analyze_bytes(encoded(Document())))
    patch = ReviewPatch(expected_version=0, candidates=[CandidateDecision(
        candidate_id=c.candidate_id, action="APPROVE" if c.scope == "document" else "REJECT",
        reason="Approve page settings for this synthetic fixture only.") for c in draft.analysis.candidates])
    reviewed = repository.change(draft.template_id, draft.revision_id, apply_review, patch)
    return repository.change(draft.template_id, draft.revision_id, publish, reviewed.version)


def test_matching_report_has_only_supported_passes(published):
    data = encoded(Document())
    result = check_bytes(published, data)
    assert result.outcome == "PASS_SUPPORTED_CHECKS"
    assert result.counts == {"PASS": 5, "FAIL": 0, "NOT_CHECKED": 0, "OUT_OF_SCOPE": 0}
    assert result.report_sha256 == sha256(data).hexdigest()
    assert result.snapshot_sha256 == published.snapshot_sha256
    assert all(i.source_path and i.evidence_ids for i in result.items)


def test_second_section_failure_is_localized(published):
    report = Document()
    section = report.add_section(WD_SECTION_START.NEW_PAGE)
    section.left_margin = Pt(30)
    result = check_bytes(published, encoded(report))
    failures = [i for i in result.items if i.status == "FAIL"]
    assert result.outcome == "FAIL" and len(failures) == 1
    assert failures[0].section_index == 1
    assert failures[0].expected["side"] == "left"
    assert failures[0].actual == {"margin_pt": 30.0}
    assert result.counts["PASS"] == 9


@pytest.mark.parametrize("change", ["missing", "malformed", "negative", "duplicate"])
def test_unresolvable_setting_abstains(published, change):
    report = Document()
    section = report.sections[0]._sectPr
    size = section.find(qn("w:pgSz"))
    if change == "missing":
        section.remove(size)
    elif change == "duplicate":
        section.append(OxmlElement("w:pgSz"))
    else:
        size.set(qn("w:w"), "595.3pt" if change == "malformed" else "-1")
    result = check_bytes(published, encoded(report))
    assert result.outcome == "INDETERMINATE"
    assert result.counts["NOT_CHECKED"] == 1
    assert next(i for i in result.items if i.status == "NOT_CHECKED").code == "UNRESOLVED_SECTION_PROPERTY"


@pytest.mark.parametrize("change", ["history", "missing_final", "wrapped", "altchunk"])
def test_ambiguous_section_layout_does_not_pass(published, change):
    report = Document()
    body = report._element.body
    if change == "history":
        report.sections[0]._sectPr.append(OxmlElement("w:sectPrChange"))
    elif change == "missing_final":
        body.remove(report.sections[0]._sectPr)
    elif change == "wrapped":
        wrapper = OxmlElement("w:sdt")
        wrapper.append(OxmlElement("w:sectPr"))
        body.insert(0, wrapper)
    else:
        body.insert(0, OxmlElement("w:altChunk"))
    result = check_bytes(published, encoded(report))
    assert result.outcome == "INDETERMINATE"
    assert result.counts["NOT_CHECKED"] == 5


@pytest.mark.parametrize("change", ["mirror", "gutter", "negative_gutter"])
def test_complex_margins_stay_unchecked(published, change):
    report = Document()
    if change == "mirror":
        report.settings.element.append(OxmlElement("w:mirrorMargins"))
    else:
        report.sections[0]._sectPr.find(qn("w:pgMar")).set(qn("w:gutter"), "100" if change == "gutter" else "-1")
    result = check_bytes(published, encoded(report))
    assert result.counts["PASS"] == 1 and result.counts["NOT_CHECKED"] == 4


def test_explicit_disabled_mirroring_does_not_block(published):
    report = Document()
    setting = OxmlElement("w:mirrorMargins")
    setting.set(qn("w:val"), "false")
    report.settings.element.append(setting)
    assert check_bytes(published, encoded(report)).outcome == "PASS_SUPPORTED_CHECKS"


def test_tolerance_is_bounded(published):
    report = Document()
    original = report.sections[0].left_margin.pt
    report.sections[0].left_margin = Pt(original + 0.1)
    assert check_bytes(published, encoded(report)).outcome == "PASS_SUPPORTED_CHECKS"
    report.sections[0].left_margin = Pt(original + 0.15)
    assert check_bytes(published, encoded(report)).outcome == "FAIL"


def test_complex_rules_keep_all_unchecked_dispositions(repository):
    draft = repository.create("Complex fixture", analyze_bytes(encoded(complex_format())))
    patch = complete_patch(draft, conditional=True)
    patch.ledger[0].action = "OUT_OF_SCOPE"
    reviewed = repository.change(draft.template_id, draft.revision_id, apply_review, patch)
    revision = repository.change(draft.template_id, draft.revision_id, publish, reviewed.version)
    result = check_bytes(revision, encoded(Document()))
    codes = {i.code for i in result.items}
    assert {"UNSUPPORTED_RULE", "UNSUPPORTED_CHAPTER_PROFILE", "DEFERRED_CANDIDATE", "REQUIREMENT_OUT_OF_SCOPE"} <= codes
    assert result.counts["OUT_OF_SCOPE"] == 1
    for rule in revision.publication.rules:
        assert any(i.rule_id == rule.rule_id for i in result.items)
    for entry in revision.publication.requirement_dispositions:
        assert any(i.item_id == "ledger:" + entry.evidence_id for i in result.items)


def test_old_snapshot_remains_usable_and_unmodified(repository, published):
    data = encoded(Document())
    before = check_bytes(published, data)
    fork = repository.fork(published.template_id, published.revision_id)
    repository.change(fork.template_id, fork.revision_id, publish, 0)
    old = repository.get(published.template_id, published.revision_id)
    assert check_bytes(old, data).model_dump() == before.model_dump()
    # Model compatibility with pre-checker compiler 0.1 snapshots.
    old.publication.compiler_version = "0.1.0"
    old.publication.report_checking_available = False
    old.snapshot_sha256 = snapshot_hash(old)
    loaded = type(old).model_validate_json(old.model_dump_json())
    assert snapshot_hash(loaded) == old.snapshot_sha256
    assert check_bytes(loaded, data).outcome == "PASS_SUPPORTED_CHECKS"


def test_api_checks_and_input_errors(repository, published):
    app.dependency_overrides[get_repository] = lambda: repository
    try:
        with TestClient(app) as client:
            base = f"/api/v2/templates/{published.template_id}/revisions/{published.revision_id}/check"
            data = encoded(Document())
            response = client.post(base, files={"file": ("report.docx", data)})
            assert response.status_code == 200 and response.json()["outcome"] == "PASS_SUPPORTED_CHECKS"
            assert client.post(base, files={"file": ("report.pdf", data)}).status_code == 415
            assert client.post(base, files={"file": ("report.docx", b"invalid")}).status_code == 422
            assert client.post(base.replace(published.template_id, "wrong"), files={"file": ("report.docx", data)}).status_code == 404
            draft = repository.fork(published.template_id, published.revision_id)
            assert client.post(base.replace(published.revision_id, draft.revision_id), files={"file": ("report.docx", data)}).status_code == 409
            assert repository.get(published.template_id, published.revision_id).model_dump() == published.model_dump()
    finally:
        app.dependency_overrides.pop(get_repository, None)


def test_corrupt_snapshot_cannot_be_checked(published):
    published.name = "unexpected change"
    with pytest.raises(ReviewError) as error:
        check_bytes(published, encoded(Document()))
    assert error.value.code == "SNAPSHOT_INTEGRITY_ERROR"


def test_matching_page_settings_do_not_hide_unsupported_rules(repository):
    draft = repository.create("Complex fixture", analyze_bytes(encoded(complex_format())))
    edited = repository.change(draft.template_id, draft.revision_id, apply_review, complete_patch(draft))
    revision = repository.change(draft.template_id, draft.revision_id, publish, edited.version)
    report = Document()
    for rule in revision.publication.rules:
        if rule.condition != "ALWAYS":
            continue
        value = rule.value
        if value.kind == "page_size":
            report.sections[0].page_width = Pt(value.width_pt)
            report.sections[0].page_height = Pt(value.height_pt)
        elif value.kind == "margin":
            setattr(report.sections[0], value.side + "_margin", Pt(value.expected_pt))
    result = check_bytes(revision, encoded(report))
    assert result.counts["FAIL"] == 0 and result.counts["PASS"] == 5
    assert result.counts["NOT_CHECKED"] > 0
    assert result.outcome == "INDETERMINATE"


def test_excessive_section_count_is_bounded(published):
    report = Document()
    for _ in range(1000):
        paragraph = report.add_paragraph()
        paragraph._p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    with pytest.raises(InvalidDocxError, match="1000-section"):
        check_bytes(published, encoded(report))

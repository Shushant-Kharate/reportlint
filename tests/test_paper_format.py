"""Property mutations must alter the relevant verdict, with scope controls."""
from io import BytesIO
import pytest
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.builtin_formats import builtin, CatalogRepository
from app.checking.service import check_bytes
from app.review.repository import RevisionRepository
from app.review.service import snapshot_hash

IEEE = 'builtin-ieee-conference-letter'
TEXT = ('This experiment measures a repeatable process and explains the observations in ordinary sentences. '
        'The researchers describe their method and compare several independent observations carefully. ') * 6


def paper():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Pt(54), Pt(72)
    section.left_margin = section.right_margin = Pt(45)
    section._sectPr.find(qn('w:cols')).set(qn('w:num'), '2')
    doc.styles['Normal'].font.name = 'Times New Roman'
    for name, size in [('Title',24), ('Heading 1',10), ('Body Text',10), ('Caption',8)]:
        doc.styles[name].font.size = Pt(size)
    doc.styles['Body Text'].paragraph_format.line_spacing = .95
    doc.styles['Body Text'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph('Measuring a repeatable process', 'Title')
    abstract = doc.add_paragraph('Abstract—This study measures a repeatable process.')
    abstract.runs[0].font.size = Pt(9)
    doc.add_paragraph('Keywords—measurement, repeatability')
    doc.add_paragraph('I. INTRODUCTION', 'Heading 1')
    doc.add_paragraph(TEXT, 'Body Text')
    doc.add_paragraph(TEXT, 'Body Text')
    doc.add_paragraph('Fig. 1. A repeatable process.', 'Caption')
    doc.add_paragraph('REFERENCES', 'Heading 1')
    doc.add_paragraph('[1] A. Writer, A fictional study.').runs[0].font.size = Pt(8)
    return doc


def check(doc):
    stream = BytesIO(); doc.save(stream)
    return check_bytes(builtin(IEEE), stream.getvalue())


def failures(result, kind):
    return [i for i in result.items if i.status == 'FAIL' and (i.expected or {}).get('kind') == kind]


def test_complete_synthetic_control_has_no_confirmed_failures():
    r = check(paper())
    assert r.counts['FAIL'] == 0
    assert r.counts['NOT_CHECKED'] > 0  # no fabricated whole-paper pass


@pytest.mark.parametrize('kind,mutate', [
    ('font_size', lambda d: setattr(d.styles['Body Text'].font, 'size', Pt(12))),
    ('line_spacing_multiple', lambda d: setattr(d.styles['Body Text'].paragraph_format, 'line_spacing', 1.5)),
    ('paper_body_alignment', lambda d: setattr(d.styles['Body Text'].paragraph_format, 'alignment', WD_ALIGN_PARAGRAPH.LEFT)),
    ('paper_abstract_size', lambda d: setattr(d.paragraphs[1].runs[0].font, 'size', Pt(14))),
    ('paper_title_size', lambda d: setattr(d.styles['Title'].font, 'size', Pt(20))),
    ('paper_heading_size', lambda d: setattr(d.styles['Heading 1'].font, 'size', Pt(16))),
    ('paper_reference_size', lambda d: setattr(d.paragraphs[-1].runs[0].font, 'size', Pt(12))),
    ('paper_heading_numbering', lambda d: setattr(d.paragraphs[3], 'text', 'CHAPTER 1 INTRODUCTION')),
    ('paper_abstract_run_in', lambda d: setattr(d.paragraphs[1], 'text', 'ABSTRACT')),
    ('paper_caption_numbering', lambda d: setattr(d.paragraphs[6], 'text', 'Fig 3.1 Architecture')),
    ('paper_keywords', lambda d: setattr(d.paragraphs[2], 'text', '')),
    ('margin', lambda d: setattr(d.sections[0], 'top_margin', Pt(65))),
    ('margin_range', lambda d: setattr(d.sections[0], 'right_margin', Pt(18))),
])
def test_mutations_detect_the_property_without_false_failures_in_control(kind, mutate):
    original = paper(); changed = paper()
    assert not failures(check(original), kind)
    mutate(changed)
    assert failures(check(changed), kind)


def test_abstract_is_not_also_checked_as_ten_point_body_text():
    doc = paper()
    doc.paragraphs[1].style = doc.styles['Body Text']
    doc.paragraphs[1].runs[0].font.size = Pt(9)
    r = check(doc)
    assert not failures(r, 'font_size')
    assert not failures(r, 'paper_abstract_size')


def test_body_size_is_not_applied_to_cover_or_after_references():
    doc = paper()
    p = doc.add_paragraph('Signature and report appendix', 'Body Text')
    p.runs[0].font.size = Pt(14)
    assert not failures(check(doc), 'font_size')


def test_nested_or_invalid_page_field_is_reviewed_but_known_field_fails():
    doc = paper()
    p = doc.sections[0].footer.paragraphs[0]
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    p._p.append(field)
    r = check(doc)
    assert len(failures(r, 'paper_no_page_numbers')) == 1
    assert failures(r, 'paper_no_page_numbers')[0].actual['page_fields'] == ['PAGE']


def test_missing_title_role_is_review_not_guessed_declaration():
    doc = paper(); doc.paragraphs[0].style = doc.styles['Normal']
    doc.add_paragraph('Declaration', 'Title')
    items = [i for i in check(doc).items if (i.expected or {}).get('kind') == 'paper_title_size']
    assert [i.status for i in items] == ['NOT_CHECKED']


def test_source_margin_variants_are_accepted_but_outside_range_fails():
    for value in [44.65, 45.0, 45.35]:
        doc = paper(); doc.sections[0].left_margin = Pt(value)
        assert not failures(check(doc), 'margin_range')
    doc.sections[0].left_margin = Pt(46)
    assert failures(check(doc), 'margin_range')


def test_historical_snapshot_remains_replayable_and_latest_is_two(tmp_path):
    catalog = CatalogRepository(RevisionRepository(tmp_path/'reviews.sqlite3'))
    latest = builtin(IEEE); old = catalog.get(IEEE, IEEE+'-v1')
    assert old.revision_number == 1 and latest.revision_number == 2
    assert snapshot_hash(old) == old.snapshot_sha256
    assert len(old.publication.rules) == 3 and len(latest.publication.rules) == 19
    assert [x['number'] for x in catalog.list_revisions(IEEE)] == [2, 1]

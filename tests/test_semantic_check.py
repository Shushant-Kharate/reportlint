from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.shared import Pt
import pytest

from app.analysis.service import analyze_bytes
from app.checking.service import check_bytes
from app.ooxml.constants import qn
from app.review.repository import RevisionRepository
from app.review.service import apply_review, publish
from tests.test_revision_review import complete_patch
from tests.test_template_analysis import encoded, complex_format


@pytest.fixture
def revision(tmp_path):
    repo = RevisionRepository(tmp_path / "review.sqlite3")
    draft = repo.create("Semantic fixture", analyze_bytes(encoded(complex_format())))
    edited = repo.change(draft.template_id, draft.revision_id, apply_review, complete_patch(draft))
    return repo.change(draft.template_id, draft.revision_id, publish, edited.version)


def report():
    doc = Document()
    body = doc.styles["Body Text"]
    body.font.name = "Times New Roman"
    body.font.size = Pt(12)
    body.paragraph_format.line_spacing = 1.5
    doc.add_heading("A fictional cover", 0)
    for title in ["Introduction", "System Design", "References"]:
        doc.add_heading(title, 1)
        doc.add_paragraph("Short body text.", style="Body Text")
    return doc


def items(revision, doc, prefix):
    return [i for i in check_bytes(revision, encoded(doc)).items if i.code.startswith(prefix)]


def test_body_and_chapters_with_no_false_cover_failures(revision):
    result = check_bytes(revision, encoded(report()))
    body = [i for i in result.items if i.code.startswith("BODY_")]
    assert len(body) == 9 and all(i.status == "PASS" for i in body)
    assert all(i.paragraph_index in {2, 4, 6} for i in body)
    assert len([i for i in result.items if i.code == "CHAPTER_PRESENT"]) == 3
    assert next(i for i in result.items if i.code.startswith("CHAPTER_ORDER")).status == "PASS"


def test_mixed_run_error_keeps_exact_location(revision):
    doc = report()
    doc.paragraphs[2].add_run(" Wrong-sized phrase.").font.size = Pt(10)
    failures = [i for i in items(revision, doc, "BODY_RUN") if i.status == "FAIL"]
    assert len(failures) == 1
    assert failures[0].paragraph_index == 2 and failures[0].run_index == 1
    assert failures[0].actual == {"value": 10.0}
    assert failures[0].source_path.endswith("w:r[2]")


def test_spacing_exact_is_not_multiple(revision):
    doc = report()
    doc.paragraphs[2].paragraph_format.line_spacing = Pt(18)
    failed = [i for i in items(revision, doc, "BODY_SPACING") if i.status == "FAIL"]
    assert len(failed) == 1 and failed[0].actual["rule"] == "exact"


def test_inherited_body_style_and_character_override(revision):
    doc = report()
    child = doc.styles.add_style("Project prose", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = doc.styles["Body Text"]
    doc.paragraphs[2].style = child
    char = doc.styles.add_style("Wrong face", WD_STYLE_TYPE.CHARACTER)
    char.font.name = "Arial"
    doc.paragraphs[2].runs[0].style = char
    failures = [i for i in items(revision, doc, "BODY_RUN") if i.status == "FAIL"]
    assert len(failures) == 1 and failures[0].actual == {"value": "Arial"}


@pytest.mark.parametrize("kind", ["theme", "non_latin", "style_cycle", "missing_style", "numbering"])
def test_unsupported_body_values_abstain(revision, kind):
    doc = report()
    para = doc.paragraphs[2]
    if kind == "theme":
        fonts = para.runs[0]._r.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn("w:asciiTheme"), "minorHAnsi")
    elif kind == "non_latin":
        para.runs[0].text = "Non-Latin: हिंदी"
    elif kind == "style_cycle":
        doc.styles["Body Text"].base_style = doc.styles["Body Text"]
    elif kind == "missing_style":
        para._p.get_or_add_pPr().get_or_add_pStyle().set(qn("w:val"), "Missing")
    else:
        para._p.get_or_add_pPr().append(OxmlElement("w:numPr"))
    result = check_bytes(revision, encoded(doc))
    assert any(i.status == "NOT_CHECKED" and (i.paragraph_index == 2 or i.code == "INCOMPLETE_BODY_SCOPE") for i in result.items)
    assert not any(i.status == "FAIL" and i.paragraph_index == 2 for i in result.items)


def test_normal_paragraph_is_not_assumed_to_be_body(revision):
    doc = report()
    doc.paragraphs[2].style = doc.styles["Normal"]
    result = check_bytes(revision, encoded(doc))
    assert any(i.code == "INCOMPLETE_BODY_SCOPE" for i in result.items)
    assert not any(i.code.startswith("BODY_") and i.paragraph_index == 2 for i in result.items)


def test_missing_duplicate_and_wrong_order_chapters(revision):
    doc = report()
    doc.paragraphs[3].text = "Unrelated chapter"
    assert any(i.code == "MISSING_CHAPTER" and i.expected["name"] == "System Design" for i in items(revision, doc, "MISSING"))
    doc = report()
    doc.add_heading("Introduction", 1)
    assert len(items(revision, doc, "DUPLICATE_CHAPTER")) == 1
    doc = report()
    doc.paragraphs[1].text, doc.paragraphs[3].text = "System Design", "Introduction"
    assert items(revision, doc, "CHAPTER_ORDER")[0].status == "FAIL"


def test_table_mention_cannot_replace_chapter(revision):
    doc = report()
    doc.paragraphs[3].text = "Unrelated chapter"
    table = doc.add_table(rows=1, cols=1)
    para = table.cell(0, 0).paragraphs[0]
    para.style = doc.styles["Heading 1"]
    para.text = "System Design"
    assert len(items(revision, doc, "MISSING_CHAPTER")) == 1


def test_contents_fields_do_not_count_as_chapters(revision):
    doc = report()
    doc.paragraphs[3].text = "Unrelated chapter"
    para = doc.add_paragraph(style="Heading 1")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    para.add_run()._r.append(begin)
    doc.add_paragraph("System Design", style="Heading 1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    doc.add_paragraph().add_run()._r.append(end)
    assert len(items(revision, doc, "MISSING_CHAPTER")) == 1


def test_unstyled_or_lower_level_title_requires_review(revision):
    for style in ("Normal", "Heading 2"):
        doc = report()
        doc.paragraphs[3].style = doc.styles[style]
        assert len(items(revision, doc, "AMBIGUOUS_CHAPTER")) == 1
        assert not items(revision, doc, "MISSING_CHAPTER")


def test_direct_outline_override_wins_over_style(revision):
    doc = report()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "9")
    doc.paragraphs[3]._p.get_or_add_pPr().append(outline)
    assert len(items(revision, doc, "AMBIGUOUS_CHAPTER")) == 1


def test_reviewed_alias_and_optional_chapter(tmp_path):
    repo = RevisionRepository(tmp_path / "review.sqlite3")
    draft = repo.create("Alias fixture", analyze_bytes(encoded(complex_format())))
    patch = complete_patch(draft)
    patch.profile.chapters[0].aliases = ["Chapter 1: Introduction"]
    patch.profile.chapters[2].required = False
    edited = repo.change(draft.template_id, draft.revision_id, apply_review, patch)
    revision = repo.change(draft.template_id, draft.revision_id, publish, edited.version)
    doc = report()
    doc.paragraphs[1].text = "  CHAPTER 1:   Introduction "
    doc.paragraphs[5].text = "An extra section"
    assert len(items(revision, doc, "CHAPTER_PRESENT")) == 2
    assert len(items(revision, doc, "OPTIONAL_CHAPTER_ABSENT")) == 1
    assert items(revision, doc, "CHAPTER_ORDER")[0].status == "PASS"


def test_custom_heading_inherits_outline_level(revision):
    doc = report()
    style = doc.styles.add_style("Project chapter", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Heading 1"]
    doc.paragraphs[3].style = style
    assert len(items(revision, doc, "CHAPTER_PRESENT")) == 3


def test_paragraph_order_preserves_nested_table_position(revision):
    doc = report()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table prose"
    doc._element.body.remove(table._tbl)
    doc._element.body.insert(3, table._tbl)
    failures_before = items(revision, doc, "CHAPTER_PRESENT")
    assert [i.paragraph_index for i in failures_before] == [1, 4, 6]


def test_body_hyperlink_runs_are_checked(revision):
    doc = report()
    link = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    props.append(size)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "Link text"
    run.append(text)
    link.append(run)
    doc.paragraphs[2]._p.append(link)
    errors = [i for i in items(revision, doc, "BODY_RUN") if i.status == "FAIL"]
    assert len(errors) == 1 and "w:hyperlink" in errors[0].source_path


def test_inline_field_does_not_silently_drop_body_coverage(revision):
    doc = report()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "REF example")
    doc.paragraphs[2]._p.append(field)
    result = check_bytes(revision, encoded(doc))
    assert any(i.code == "INCOMPLETE_BODY_SCOPE" and i.actual["unresolved_paragraphs"] == 1 for i in result.items)
    assert not any(i.code.startswith("BODY_") and i.paragraph_index == 2 for i in result.items)

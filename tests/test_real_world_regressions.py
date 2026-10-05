"""Fictional examples of problems discovered in the private real-file audit."""
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt

from app.analysis.service import analyze_bytes
from app.checking.role_review import preview_bytes, ReportRoleReview, RoleDecision
from app.checking.service import check_bytes
from app.ooxml.constants import qn
from app.review.repository import RevisionRepository
from app.review.service import apply_review, publish
from tests.test_revision_review import complete_patch
from tests.test_semantic_check import revision, report
from tests.test_template_analysis import encoded, complex_format


def test_numbered_lower_level_chapters_and_punctuation(revision):
    doc = report()
    doc.paragraphs[1].style = 'Heading 2'
    doc.paragraphs[1].text = 'CHAPTER 1 INTRODUCTION'
    doc.paragraphs[3].style = 'Heading 2'
    doc.paragraphs[3].text = 'Chapter 2: System Design'
    doc.paragraphs[5].text = 'REFERENCES: '
    result = check_bytes(revision, encoded(doc))
    assert len([i for i in result.items if i.code == 'CHAPTER_PRESENT']) == 3
    assert not any(i.code in {'MISSING_CHAPTER','AMBIGUOUS_CHAPTER'} for i in result.items)


def test_split_banner_and_subheading_do_not_make_duplicate(revision):
    doc = report()
    doc.paragraphs[3].style = 'Heading 2'
    doc.paragraphs[3].text = 'CHAPTER 2'
    title = doc.paragraphs[4].insert_paragraph_before('System Design', style='Heading 2')
    doc.paragraphs[5].insert_paragraph_before('System Design', style='Heading 3')
    result = check_bytes(revision, encoded(doc))
    assert len([i for i in result.items if i.code == 'CHAPTER_PRESENT']) == 3


def test_missing_spacing_is_default_single_not_unknown(revision):
    doc = report()
    doc.styles['Body Text'].paragraph_format.line_spacing = None
    defaults = doc.styles.element.find(qn('w:docDefaults'))
    for spacing in defaults.findall('.//' + qn('w:spacing')):
        spacing.getparent().remove(spacing)
    result = check_bytes(revision, encoded(doc))
    spacing = [i for i in result.items if i.code == 'BODY_SPACING_FAIL']
    assert len(spacing) == 2  # References are a separate scope needing review.
    assert all(i.actual['multiplier'] == 1 and i.actual['origin'] == 'default' for i in spacing)


def test_spacing_tolerance_stays_explicit_not_relaxed_to_fit_report(revision):
    doc = report()
    doc.paragraphs[2].paragraph_format.line_spacing = 362 / 240
    result = check_bytes(revision, encoded(doc))
    failure = next(i for i in result.items if i.code == 'BODY_SPACING_FAIL')
    assert '1/240' in failure.message


def test_cover_body_style_requires_scope_review_and_can_be_corrected(revision):
    doc = report()
    doc.paragraphs[0].style = 'Body Text'
    doc.paragraphs[0].runs[0].font.size = Pt(20)
    data = encoded(doc)
    result = check_bytes(revision, data)
    assert any(i.code == 'BODY_SCOPE_REVIEW' and i.paragraph_index == 0 for i in result.items)
    assert not any(i.status == 'FAIL' and i.paragraph_index == 0 for i in result.items)
    preview = preview_bytes(revision, data)
    p = preview.paragraphs[0]
    assert p.reviewable
    review = ReportRoleReview(report_sha256=preview.report_sha256, snapshot_sha256=preview.snapshot_sha256,
        decisions=[RoleDecision(paragraph_index=0, source_path=p.source_path, role='EXCLUDE', reason='This is cover lettering, not body prose.')])
    result = check_bytes(revision, data, review)
    assert not any(i.code == 'BODY_SCOPE_REVIEW' and i.paragraph_index == 0 for i in result.items)


def test_existing_heading_can_be_explicitly_remapped(revision):
    doc = report()
    doc.paragraphs[3].text = 'Detailed Architecture'
    data = encoded(doc)
    preview = preview_bytes(revision, data)
    p = preview.paragraphs[3]
    assert p.reviewable
    review = ReportRoleReview(report_sha256=preview.report_sha256, snapshot_sha256=preview.snapshot_sha256,
        decisions=[RoleDecision(paragraph_index=3, source_path=p.source_path, role='CHAPTER', chapter_index=1, reason='Reviewer verified this is the system design chapter.')])
    assert len([i for i in check_bytes(revision, data, review).items if i.code == 'CHAPTER_PRESENT']) == 3


def test_new_scoped_rules_require_review_and_find_errors(tmp_path):
    template = complex_format()
    template.add_paragraph('Chapter number and title shall be printed in font size (18pt), using both upper and lower case.')
    template.add_paragraph('The 5 word abstract shall summarize the report. The Abstract shall include key words.')
    analysis = analyze_bytes(encoded(template))
    assert {c.value.kind for c in analysis.candidates} >= {'abstract_word_count','abstract_keywords','chapter_case','chapter_font_size'}
    repo = RevisionRepository(tmp_path/'scoped.sqlite3')
    draft = repo.create('Fictional scoped template', analysis)
    patch = complete_patch(draft)
    edited = repo.change(draft.template_id, draft.revision_id, apply_review, patch)
    published = repo.change(draft.template_id, draft.revision_id, publish, edited.version)
    doc = report()
    doc.paragraphs[1].insert_paragraph_before('Abstract', style='Heading 1')
    abstract = doc.paragraphs[2].insert_paragraph_before('This summary has three extra words.', style='Body Text')
    for p in doc.paragraphs:
        if p.text == 'Introduction':
            p.text = 'CHAPTER 1 INTRODUCTION'
            p.style = 'Heading 2'
            for r in p.runs: r.font.size = Pt(16)
    result = check_bytes(published, encoded(doc))
    codes={i.code for i in result.items}
    assert codes >= {'ABSTRACT_WORD_COUNT_FAIL','ABSTRACT_KEYWORDS_FAIL','CHAPTER_FONT_SIZE_FAIL','CHAPTER_CASE_FAIL'}
    abstract.text = 'This summary has five words.'
    for p in doc.paragraphs:
        if p.text == 'CHAPTER 1 INTRODUCTION':
            p.insert_paragraph_before('Keywords: design, example', style='Body Text')
            p.text = 'Chapter 1 Introduction'
            for r in p.runs: r.font.size = Pt(18)
    codes={i.code for i in check_bytes(published, encoded(doc)).items}
    assert codes >= {'ABSTRACT_WORD_COUNT_PASS','ABSTRACT_KEYWORDS_PASS','CHAPTER_FONT_SIZE_PASS','CHAPTER_CASE_PASS'}


def test_abstract_without_end_boundary_does_not_guess(tmp_path):
    from app.checking.scoped import scoped_checks
    from app.checking.roles import inventory
    from app.models.template_analysis import AbstractWordsValue
    from app.review.models import CompiledRule
    from app.ooxml.docx_loader import DocxPackage
    doc = Document(); doc.add_heading('Abstract', 1); doc.add_paragraph('Short text')
    path = tmp_path/'report.docx'; doc.save(path)
    package = DocxPackage.load(str(path)); paragraphs, registry, uncertain = inventory(package)
    rule = CompiledRule(rule_id='abstract', candidate_ids=[], evidence_ids=[], scope='document', condition='ALWAYS', value=AbstractWordsValue(expected_words=500), readiness='PROPERTY_CONTRACT')
    assert scoped_checks(rule, package, paragraphs, registry, uncertain)[0].status == 'NOT_CHECKED'


def test_negative_scoped_instructions_are_not_positive_rules():
    doc = Document()
    doc.add_paragraph('Guidelines for project report format')
    doc.add_paragraph('Chapter number and title must not be printed at 18 pt.')
    doc.add_paragraph('The Abstract shall not include keywords.')
    result = analyze_bytes(encoded(doc))
    assert not any(c.value.kind in {'chapter_font_size', 'abstract_keywords'} for c in result.candidates)

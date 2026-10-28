"""Create fictional inputs for verify_automatic_setup.cjs, outside the repository.

Usage: python -m scripts.create_automatic_fixtures OUTPUT_DIRECTORY
Existing fixture files are never overwritten.
"""
import argparse
from pathlib import Path
from uuid import uuid4
from docx import Document
from docx.shared import Inches, Pt


def create(directory):
    directory.mkdir(parents=True, exist_ok=True)
    if any((directory/name).exists() for name in ('unseen-format.docx','report.docx')):
        raise FileExistsError('Choose another output directory; fixture files already exist.')
    format_doc=Document()
    format_doc.core_properties.identifier=uuid4().hex
    for line in ['Use Letter paper.','All margins must be 1 inch.',
                 'Body text must use 12 pt with 1.5 line spacing in one column.',
                 'Body text must use Arial or Times New Roman.']:
        format_doc.add_paragraph(line)
    format_doc.save(directory/'unseen-format.docx')
    report=Document();section=report.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
    report.styles['Body Text'].font.name='Times New Roman'
    report.styles['Body Text'].font.size=Pt(12)
    report.styles['Body Text'].paragraph_format.line_spacing=1.5
    report.add_paragraph('Fictional report for automatic setup','Title')
    text=('This fictional study measures a repeatable process and records its observations in ordinary sentences. '
          'The example compares an expected value with a recorded value and explains how the difference is calculated. ')*4
    for _ in range(2):report.add_paragraph(text,'Body Text')
    report.save(directory/'report.docx')
    print('Created two fictional browser fixtures in',directory.resolve())


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output_directory',type=Path)
    create(parser.parse_args().output_directory)

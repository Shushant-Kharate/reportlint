"""Create three small, fictional DOCX files for the ReportLint walkthrough."""

from pathlib import Path

from docx import Document
from docx.shared import Pt


SECTIONS = ("Introduction", "Methods", "Results", "Conclusion")
TEXT = {
    "Introduction": "This sample report describes a fictional study of how students organize their project notes and prepare their final documents.",
    "Methods": "The fictional team compared four sample note taking workflows and recorded the time needed to find each reference again.",
    "Results": "In this invented example, a consistent naming scheme made references easier to find and helped the team prepare its report.",
    "Conclusion": "The sample concludes that a clear document structure and consistent formatting make a report easier to review.",
}


def make_document(path: Path, *, bad_format: bool = False, template: bool = False) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(595.3)
    section.page_height = Pt(841.9)
    for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, edge, Pt(50 if bad_format else 72))

    title = doc.add_paragraph(style="Title")
    title.add_run("Sample Project Report Format" if template else "Sample Project Report")

    for name in SECTIONS:
        if bad_format and name == "Methods":
            continue  # An intentional missing section in addition to formatting errors.
        doc.add_paragraph(name, style="Heading 1")
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.0 if bad_format else 1.5
        run = paragraph.add_run(TEXT[name])
        run.font.name = "Arial" if bad_format else "Times New Roman"
        run.font.size = Pt(14 if bad_format else 12)

    doc.save(path)


def create_demo_files(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    make_document(directory / "format_template.docx", template=True)
    make_document(directory / "correct_report.docx")
    make_document(directory / "formatting_errors_report.docx", bad_format=True)


if __name__ == "__main__":
    create_demo_files(Path(__file__).resolve().parent.parent / "demo")
